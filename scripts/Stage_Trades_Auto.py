"""
Stage Trades Auto — Build Live_Trade_Info from Latest Earnings

Modes (Live_Trade_Info Trade_Mode!C3):
  - Single Day / weekday (legacy): column O flags
      "T" → Daily_Trades; "1"–"5" → Monday–Friday sheets.
  - Full_Auto: ignore O; use columns M + K + Y (+ AC). Stage from the first
      trade-row C in column Q downward; if none, fall back to first M1.
      Opens Latest Earnings once via xlwings, stages Live_Trade_Info, runs
      price flagging on the same book, then saves Earnings once at the end.

Live_Trade_Info per staged row:
  A ticker, B Direction (ToS_Open) from Y, C Share Size (Z),
  D ToS_Open_Exit (AB), E ToS Exit (AA), F Direction (ToS) from AC.
"""

import os
import sys
import warnings
from collections import defaultdict
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment
from openpyxl.worksheet.worksheet import Worksheet

warnings.filterwarnings("ignore", message=".*Unknown extension.*", category=UserWarning)
warnings.filterwarnings("ignore", message=".*Conditional Formatting extension.*", category=UserWarning)

from full_auto_utils import (
    MODE_FULL_AUTO,
    is_full_auto_mode,
    parse_earnings_date,
    parse_m_dow,
    read_trade_mode_value,
    resolve_trade_date,
    sheet_name_for_trade_date,
)

# ---------------------------------------------------------------------------
# Paths
# ---------------------------------------------------------------------------
_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
_BASE_DIR = os.path.dirname(_SCRIPT_DIR)

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------
SOURCE_FILE = os.path.join(_BASE_DIR, "! -- Latest Earnings Document.xlsx")
SOURCE_SHEET = "Trades"
HEADER_ROW = 3  # Data starts the row after headers (e.g. row 4)
COL_FLAG = "O"
COL_TICKER = "A"
COL_K = "K"
COL_M = "M"
COL_Q = "Q"
COL_DIRECTION_IBKR = "Y"   # Live_Trade_Info B (IBKR / ToS2)
COL_DIRECTION_TOS = "AC"   # Live_Trade_Info F (ToS primary)
COL_SIZE = "Z"
COL_IBKR_EXIT = "AB"   # Source for Live_Trade_Info column D
COL_TOS_EXIT = "AA"    # Source for Live_Trade_Info column E

OUTPUT_FILE = os.path.join(_BASE_DIR, "Live_Trade_Info.xlsx")
DAILY_SHEET = "Daily_Trades"
WEEKDAY_SHEETS = ("Monday", "Tuesday", "Wednesday", "Thursday", "Friday")
HEADER_B1 = "Direction (ToS_Open)"
HEADER_D1 = "ToS_Open_Exit"
HEADER_E1 = "ToS Exit"
HEADER_F1 = "Direction (ToS)"

# Column O: 1 -> Monday, 2 -> Tuesday, 3 -> Wednesday, 4 -> Thursday, 5 -> Friday
_FLAG_TO_WEEKDAY = {1: "Monday", 2: "Tuesday", 3: "Wednesday", 4: "Thursday", 5: "Friday"}


def _parse_flag(cell_value: Any) -> Optional[Any]:
    """
    Parse column O. Returns "T", 1, 2, 3, 4, 5, or None (skip).
    Accepts string or int for numbers (e.g. "1" or 1 -> 1). 5 = Friday.
    """
    if cell_value is None:
        return None
    s = str(cell_value).strip().upper()
    if s == "T":
        return "T"
    try:
        n = int(cell_value) if isinstance(cell_value, (int, float)) else int(s)
        if 1 <= n <= 5:
            return n
    except (ValueError, TypeError):
        pass
    return None


def _normalize_ticker(cell_value: Any) -> Optional[str]:
    if cell_value is None:
        return None
    return str(cell_value).strip().upper()


def _normalize_direction(cell_value: Any) -> Optional[str]:
    if cell_value is None:
        return None
    return str(cell_value).strip().lower()


def _cell_to_str(cell_value: Any) -> str:
    """Return cell value as string for exit-type columns; empty string if None."""
    if cell_value is None:
        return ""
    return str(cell_value).strip()


# Trade tuple: (ticker, direction_ibkr, size, ibkr_exit, tos_exit, direction_tos)
TradeRow = Tuple[str, str, Any, str, str, str]


def _write_trades_to_sheet(ws: Worksheet, trades: List[TradeRow]) -> None:
    """Write header row and trade rows to sheet (same format as Daily_Trades)."""
    ws["A1"] = "Ticker"
    ws["B1"] = HEADER_B1
    ws["C1"] = "Share Size"
    ws["D1"] = HEADER_D1
    ws["E1"] = HEADER_E1
    ws["F1"] = HEADER_F1
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    for ticker, direction_ibkr, size, ibkr_exit, tos_exit, direction_tos in trades:
        next_row = ws.max_row + 1
        ws.cell(row=next_row, column=1, value=ticker)
        ws.cell(row=next_row, column=2, value=direction_ibkr)
        ws.cell(row=next_row, column=3, value=size)
        ws.cell(row=next_row, column=4, value=ibkr_exit or None)
        ws.cell(row=next_row, column=5, value=tos_exit or None)
        ws.cell(row=next_row, column=6, value=direction_tos)
    left_align = Alignment(horizontal="left")
    for col in range(1, 7):
        ws.cell(row=1, column=col).alignment = left_align
    for row in range(2, ws.max_row + 1):
        for col in range(1, 7):
            ws.cell(row=row, column=col).alignment = left_align


def _is_q_c_flag(cell_value: Any) -> bool:
    if cell_value is None:
        return False
    return str(cell_value).strip().upper() == "C"


def _is_q_m1_flag(cell_value: Any) -> bool:
    if cell_value is None:
        return False
    return str(cell_value).strip().upper() == "M1"


def _xw_cell(sheet, col: str, row: int) -> Any:
    """Read one cell from an xlwings sheet."""
    return sheet.range(f"{col}{row}").value


def _xw_max_row(sheet) -> int:
    """Last used row on an xlwings sheet (fallback keeps a wide scan)."""
    try:
        used = sheet.used_range
        if used is not None:
            return int(used.last_cell.row)
    except Exception:
        pass
    return 5000


def _first_trade_row_c_flag_xw(sheet, max_row: int) -> Optional[int]:
    """
    First Q C on a real trade row (column M has a day flag).
    Skips terminator-only C on blank-M rows below an M1 block.
    """
    for row in range(HEADER_ROW + 1, max_row + 1):
        if not _is_q_c_flag(_xw_cell(sheet, COL_Q, row)):
            continue
        if parse_m_dow(_xw_cell(sheet, COL_M, row)) is None:
            continue
        return row
    return None


def _first_m1_flag_row_xw(sheet, max_row: int) -> Optional[int]:
    """First Q M1 flag (fallback staging start on M1-only days)."""
    for row in range(HEADER_ROW + 1, max_row + 1):
        if _is_q_m1_flag(_xw_cell(sheet, COL_Q, row)):
            return row
    return None


def _resolve_full_auto_stage_start_xw(sheet, max_row: int) -> Optional[tuple]:
    """
    Return (start_row, reason) for Full_Auto staging.
    Prefer first trade-row C; else first M1; else None.
    """
    c_row = _first_trade_row_c_flag_xw(sheet, max_row)
    if c_row is not None:
        return c_row, f"trade-row C at {c_row}"
    m1_row = _first_m1_flag_row_xw(sheet, max_row)
    if m1_row is not None:
        return m1_row, f"M1 fallback at {m1_row} (no trade-row C)"
    return None


def _write_live_trade_info_sheets(by_sheet: Dict[str, List[TradeRow]]) -> None:
    """Write staged trades to Live_Trade_Info (separate from Earnings save)."""
    if not os.path.exists(OUTPUT_FILE):
        wb_output = Workbook()
        wb_output.remove(wb_output.active)
    else:
        wb_output = load_workbook(OUTPUT_FILE)

    all_tickers: List[str] = []
    for sheet_name in sorted(by_sheet.keys()):
        trades = by_sheet[sheet_name]
        if sheet_name not in wb_output.sheetnames:
            wb_output.create_sheet(sheet_name)
        _write_trades_to_sheet(wb_output[sheet_name], trades)
        print(f"Wrote {len(trades)} trade(s) to sheet '{sheet_name}'.")
        all_tickers.extend(t[0] for t in trades)

    wb_output.save(OUTPUT_FILE)
    try:
        wb_output.close()
    except Exception:
        pass
    print(f"Full_Auto staging complete. Tickers: {', '.join(all_tickers)}.")


def _collect_full_auto_trades_xw(sheet) -> Dict[str, List[TradeRow]]:
    """Stage M+K+Y+AC from first trade-row C (or M1 fallback) downward."""
    max_row = _xw_max_row(sheet)
    by_sheet: Dict[str, List[TradeRow]] = defaultdict(list)

    start_info = _resolve_full_auto_stage_start_xw(sheet, max_row)
    if start_info is None:
        print(
            "Full_Auto: no trade-row C and no M1 in column Q; staging nothing."
        )
        return by_sheet

    start_row, reason = start_info
    print(f"Full_Auto: staging from {reason} downward.")
    for row in range(start_row, max_row + 1):
        m_token = parse_m_dow(_xw_cell(sheet, COL_M, row))
        k_date = parse_earnings_date(_xw_cell(sheet, COL_K, row))
        direction_ibkr = _normalize_direction(
            _xw_cell(sheet, COL_DIRECTION_IBKR, row)
        )
        direction_tos = _normalize_direction(
            _xw_cell(sheet, COL_DIRECTION_TOS, row)
        )
        ticker = _normalize_ticker(_xw_cell(sheet, COL_TICKER, row))
        size = _xw_cell(sheet, COL_SIZE, row)
        ibkr_exit = _cell_to_str(_xw_cell(sheet, COL_IBKR_EXIT, row))
        tos_exit = _cell_to_str(_xw_cell(sheet, COL_TOS_EXIT, row))

        if not m_token or k_date is None:
            continue
        if not ticker:
            print(f"Row {row}: missing ticker; skipping.")
            continue
        if direction_ibkr not in ("long", "short"):
            continue
        if direction_tos not in ("long", "short"):
            print(
                f"Row {row}: missing/invalid ToS direction (AC) for ticker "
                f"{ticker}; skipping."
            )
            continue
        if size is None or str(size).strip() == "":
            print(f"Row {row}: missing share size for ticker {ticker}; skipping.")
            continue

        trade_date = resolve_trade_date(m_token, k_date)
        sheet_name = sheet_name_for_trade_date(trade_date, m_token)
        by_sheet[sheet_name].append(
            (ticker, direction_ibkr, size, ibkr_exit, tos_exit, direction_tos)
        )
    return by_sheet


def _stage_full_auto() -> None:
    """
    Full_Auto: one xlwings open of Latest Earnings — stage Live_Trade_Info,
    flag P/Q on the same book, save Earnings once, then release.
    """
    from earnings_workbook_utils import (
        nudge_onedrive_sync,
        open_or_attach_earnings_workbook,
        release_earnings_workbook,
    )
    from flag_prices_full_auto import flag_prices_full_auto

    app = None
    wb = None
    owned_app = False
    owned_book = False
    did_save_earnings = False
    try:
        app, wb, owned_app, owned_book = open_or_attach_earnings_workbook(SOURCE_FILE)
        try:
            sheet = wb.sheets[SOURCE_SHEET]
        except Exception:
            print(f"Worksheet '{SOURCE_SHEET}' not found in {SOURCE_FILE}.")
            return

        by_sheet = _collect_full_auto_trades_xw(sheet)
        if not by_sheet:
            print("Full_Auto: no eligible trades to stage (from C/M1 start downward).")
        else:
            _write_live_trade_info_sheets(by_sheet)

        print("Full_Auto: running embedded price flagging (same Excel session)...")
        needs_save = flag_prices_full_auto(
            earnings_file=SOURCE_FILE,
            sheet=sheet,
            wb=wb,
            save=False,
        )
        if needs_save:
            try:
                wb.save()
            except Exception:
                try:
                    wb.save()
                except Exception as e:
                    print(f"Full_Auto: failed to save Latest Earnings after flagging: {e}")
                    sys.exit(1)
            did_save_earnings = True
            print("Full_Auto: Latest Earnings saved once after flagging.")
        else:
            print("Full_Auto: no Earnings flag changes; skip Earnings save.")
    except Exception as e:
        print(f"Full_Auto staging/flagging failed: {e}")
        sys.exit(1)
    finally:
        release_earnings_workbook(app, wb, owned_app=owned_app, owned_book=owned_book)
        # After headless Excel quits, bump mtime so OneDrive notices the COM save.
        if did_save_earnings:
            nudge_onedrive_sync(SOURCE_FILE)


def _stage_legacy(ws_source) -> None:
    """Legacy column-O staging (Single Day / weekday)."""
    max_row = ws_source.max_row
    start_row = HEADER_ROW + 1

    t_rows: List[TradeRow] = []
    by_weekday: Dict[str, List[TradeRow]] = {
        "Monday": [], "Tuesday": [], "Wednesday": [], "Thursday": [], "Friday": []
    }

    for row in range(start_row, max_row + 1):
        flag_val = _parse_flag(ws_source[f"{COL_FLAG}{row}"].value)
        if flag_val is None:
            continue

        raw_ticker = ws_source[f"{COL_TICKER}{row}"].value
        raw_direction_ibkr = ws_source[f"{COL_DIRECTION_IBKR}{row}"].value
        raw_direction_tos = ws_source[f"{COL_DIRECTION_TOS}{row}"].value
        raw_size = ws_source[f"{COL_SIZE}{row}"].value
        raw_ibkr_exit = ws_source[f"{COL_IBKR_EXIT}{row}"].value
        raw_tos_exit = ws_source[f"{COL_TOS_EXIT}{row}"].value

        ticker = _normalize_ticker(raw_ticker)
        direction_ibkr = _normalize_direction(raw_direction_ibkr)
        direction_tos = _normalize_direction(raw_direction_tos)
        size = raw_size
        ibkr_exit = _cell_to_str(raw_ibkr_exit)
        tos_exit = _cell_to_str(raw_tos_exit)

        if not ticker:
            print(f"Row {row}: missing ticker; skipping.")
            continue
        if not direction_ibkr:
            print(
                f"Row {row}: missing IBKR direction (Y) for ticker {ticker}; skipping."
            )
            continue
        if not direction_tos:
            print(
                f"Row {row}: missing ToS direction (AC) for ticker {ticker}; skipping."
            )
            continue
        if size is None or str(size).strip() == "":
            print(f"Row {row}: missing share size for ticker {ticker}; skipping.")
            continue

        row_data: TradeRow = (
            ticker, direction_ibkr, size, ibkr_exit, tos_exit, direction_tos
        )
        if flag_val == "T":
            t_rows.append(row_data)
        else:
            day_name = _FLAG_TO_WEEKDAY[flag_val]
            by_weekday[day_name].append(row_data)

    has_weekday = any(by_weekday[d] for d in WEEKDAY_SHEETS)
    total_trades = len(t_rows) + sum(len(by_weekday[d]) for d in WEEKDAY_SHEETS)

    if total_trades == 0:
        print("No valid trades found (no rows with 'T' or 1/2/3/4/5 in column O had ticker, direction, and share size).")
        if os.path.exists(OUTPUT_FILE):
            try:
                wb_out = load_workbook(OUTPUT_FILE)
                if DAILY_SHEET in wb_out.sheetnames:
                    ws_out = wb_out[DAILY_SHEET]
                    if ws_out.max_row > 1:
                        ws_out.delete_rows(2, ws_out.max_row - 1)
                    wb_out.save(OUTPUT_FILE)
                    print("Cleared existing trades in Live_Trade_Info (Daily_Trades).")
            except PermissionError:
                print("Could not update Live_Trade_Info (file may be open). Please close it.")
                sys.exit(1)
        sys.exit(0)

    if has_weekday:
        print("No clear single day range, writing weekday trades.")
        if not os.path.exists(OUTPUT_FILE):
            wb_output = Workbook()
            wb_output.remove(wb_output.active)
        else:
            wb_output = load_workbook(OUTPUT_FILE)
        for day_name in WEEKDAY_SHEETS:
            trades_for_day = by_weekday[day_name]
            if day_name not in wb_output.sheetnames:
                wb_output.create_sheet(day_name)
            ws = wb_output[day_name]
            _write_trades_to_sheet(ws, trades_for_day)
            if trades_for_day:
                print(f"Wrote {len(trades_for_day)} trade(s) to sheet '{day_name}'.")
        wb_output.save(OUTPUT_FILE)
        all_tickers = []
        for day_name in WEEKDAY_SHEETS:
            for t in by_weekday[day_name]:
                all_tickers.append(t[0])
        print(f"Tickers: {', '.join(all_tickers)}.")
    else:
        print(f"Collected {len(t_rows)} trade(s) from Earnings (column O = T).")
        if os.path.exists(OUTPUT_FILE):
            wb_output = load_workbook(OUTPUT_FILE)
            ws_output = wb_output[DAILY_SHEET] if DAILY_SHEET in wb_output.sheetnames else wb_output.active
        else:
            wb_output = Workbook()
            ws_output = wb_output.active
        ws_output.title = DAILY_SHEET
        _write_trades_to_sheet(ws_output, t_rows)
        wb_output.save(OUTPUT_FILE)
        print(f"Wrote {len(t_rows)} trade(s) to '{OUTPUT_FILE}' (sheet '{DAILY_SHEET}').")
        print(f"Tickers: {', '.join(t[0] for t in t_rows)}.")


def main() -> None:
    if not os.path.exists(SOURCE_FILE):
        print(f"Source Earnings file not found: {SOURCE_FILE}")
        return

    mode_raw = read_trade_mode_value(OUTPUT_FILE) if os.path.exists(OUTPUT_FILE) else None
    if is_full_auto_mode(mode_raw):
        print(
            f"Trade_Mode is {MODE_FULL_AUTO}: staging from M+K+Y (ignore O); "
            f"xlwings single-session Earnings save."
        )
        try:
            _stage_full_auto()
        except PermissionError:
            print(
                "Could not update Live_Trade_Info or Latest Earnings "
                "(file may be open). Please close it."
            )
            sys.exit(1)
        return

    try:
        wb_source = load_workbook(SOURCE_FILE, data_only=True)
    except PermissionError:
        print("Please close Latest Earnings Document.")
        return

    try:
        ws_source = wb_source[SOURCE_SHEET]
    except KeyError:
        print(f"Worksheet '{SOURCE_SHEET}' not found in {SOURCE_FILE}.")
        return

    try:
        _stage_legacy(ws_source)
    finally:
        try:
            wb_source.close()
        except Exception:
            pass


if __name__ == "__main__":
    main()
