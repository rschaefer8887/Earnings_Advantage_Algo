"""
Full_Auto helpers: trade dates from columns M+K, sheet names, trading-day math.

Trade_Mode!C3 = Full_Auto stages onto date sheets like 9_10_Th and drives
order sheet selection / price flagging (see flag_prices_full_auto.py).
"""

from __future__ import annotations

import os
import re
from datetime import date, datetime, timedelta
from typing import Any, Optional, Set

TRADE_MODE_SHEET = "Trade_Mode"
TRADE_MODE_CELL = "C3"
MODE_SINGLE_DAY = "Single Day"
MODE_FULL_AUTO = "Full_Auto"

# Column M tokens -> Python weekday (Monday=0 ... Friday=4)
_M_TO_WEEKDAY = {
    "M": 0,
    "TU": 1,
    "W": 2,
    "TH": 3,
    "F": 4,
}
_WEEKDAY_TO_M = {0: "M", 1: "Tu", 2: "W", 3: "Th", 4: "F"}

DATE_SHEET_RE = re.compile(r"^(\d{1,2})_(\d{1,2})_(M|Tu|W|Th|F)$")

# NYSE holidays (observed) for current automation window — extend annually.
_NYSE_HOLIDAYS: Set[date] = {
    # 2025
    date(2025, 1, 1),
    date(2025, 1, 20),
    date(2025, 2, 17),
    date(2025, 4, 18),
    date(2025, 5, 26),
    date(2025, 6, 19),
    date(2025, 7, 4),
    date(2025, 9, 1),
    date(2025, 11, 27),
    date(2025, 12, 25),
    # 2026
    date(2026, 1, 1),
    date(2026, 1, 19),
    date(2026, 2, 16),
    date(2026, 4, 3),
    date(2026, 5, 25),
    date(2026, 6, 19),
    date(2026, 7, 3),  # Independence Day observed
    date(2026, 9, 7),
    date(2026, 11, 26),
    date(2026, 12, 25),
    # 2027
    date(2027, 1, 1),
    date(2027, 1, 18),
    date(2027, 2, 15),
    date(2027, 3, 26),
    date(2027, 5, 31),
    date(2027, 6, 18),  # Juneteenth observed
    date(2027, 7, 5),  # Independence Day observed
    date(2027, 9, 6),
    date(2027, 11, 25),
    date(2027, 12, 24),  # Christmas observed
}


def parse_m_dow(cell_value: Any) -> Optional[str]:
    """Return canonical M token (M/Tu/W/Th/F) or None."""
    if cell_value is None:
        return None
    s = str(cell_value).strip()
    if not s:
        return None
    key = s.upper().replace(".", "")
    # Accept TU / TU. / Tuesday-ish short forms
    if key in ("M", "MON"):
        return "M"
    if key in ("TU", "TUE", "TUES"):
        return "Tu"
    if key in ("W", "WED"):
        return "W"
    if key in ("TH", "THU", "THUR", "THURS"):
        return "Th"
    if key in ("F", "FRI"):
        return "F"
    return None


def parse_earnings_date(cell_value: Any) -> Optional[date]:
    """Parse column K earnings date from Excel date / datetime / string."""
    if cell_value is None:
        return None
    if isinstance(cell_value, datetime):
        return cell_value.date()
    if isinstance(cell_value, date):
        return cell_value
    s = str(cell_value).strip()
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def m_to_weekday(m_token: str) -> int:
    return _M_TO_WEEKDAY[m_token.upper() if m_token.upper() != "TU" else "TU"]


def dow_abbrev_for_date(d: date) -> str:
    """Weekday token matching column M style."""
    if d.weekday() > 4:
        raise ValueError(f"{d} is not a weekday")
    return _WEEKDAY_TO_M[d.weekday()]


def resolve_trade_date(m_token: str, earnings_date: date) -> date:
    """
    Most recent calendar date on or before earnings_date whose weekday matches M.
    """
    target_wd = m_to_weekday(m_token)
    d = earnings_date
    for _ in range(8):
        if d.weekday() == target_wd:
            return d
        d -= timedelta(days=1)
    raise ValueError(f"Could not resolve trade date for M={m_token} K={earnings_date}")


def sheet_name_for_trade_date(trade_date: date, m_token: str) -> str:
    return f"{trade_date.month}_{trade_date.day}_{m_token}"


def parse_date_sheet_name(name: str) -> Optional[tuple[date, str]]:
    """Return (date, m_token) for names like 9_10_Th, else None."""
    m = DATE_SHEET_RE.match(name.strip())
    if not m:
        return None
    month, day, token = int(m.group(1)), int(m.group(2)), m.group(3)
    # Year: prefer year of trade_date matching token near today; use current year then adjacent
    today = date.today()
    for year in (today.year, today.year - 1, today.year + 1):
        try:
            d = date(year, month, day)
        except ValueError:
            continue
        if dow_abbrev_for_date(d) == token:
            return d, token
    try:
        d = date(today.year, month, day)
        return d, token
    except ValueError:
        return None


def is_trading_day(d: date) -> bool:
    if d.weekday() >= 5:
        return False
    return d not in _NYSE_HOLIDAYS


def add_trading_days(d: date, n: int) -> date:
    """Add n trading days (n can be negative). n=0 returns d if trading else next rules unused."""
    if n == 0:
        return d
    step = 1 if n > 0 else -1
    remaining = abs(n)
    cur = d
    while remaining:
        cur += timedelta(days=step)
        if is_trading_day(cur):
            remaining -= 1
    return cur


def cell_has_value(cell_value: Any) -> bool:
    if cell_value is None:
        return False
    if isinstance(cell_value, str) and cell_value.strip() == "":
        return False
    return True


def read_trade_mode_value(file_path: str) -> Optional[str]:
    """Read Trade_Mode!C3; None if missing/unreadable."""
    if not file_path or not os.path.exists(file_path):
        return None
    try:
        from openpyxl import load_workbook

        wb = load_workbook(file_path, data_only=True)
        try:
            if TRADE_MODE_SHEET not in wb.sheetnames:
                return None
            return wb[TRADE_MODE_SHEET][TRADE_MODE_CELL].value
        finally:
            wb.close()
    except Exception:
        return None


def normalize_trade_mode(value: Any) -> Optional[str]:
    if value is None:
        return None
    s = str(value).strip()
    if not s:
        return None
    low = s.lower()
    if low == MODE_SINGLE_DAY.lower():
        return MODE_SINGLE_DAY
    if low == MODE_FULL_AUTO.lower():
        return MODE_FULL_AUTO
    return s


def is_full_auto_mode(value: Any) -> bool:
    return normalize_trade_mode(value) == MODE_FULL_AUTO


def full_auto_sheet_name_for_today(today: Optional[date] = None) -> Optional[str]:
    """Sheet name for today's Full_Auto orders, or None if not a trading day."""
    d = today or date.today()
    if not is_trading_day(d):
        return None
    return sheet_name_for_trade_date(d, dow_abbrev_for_date(d))
